"""Independently validate every final row against original sources."""
import argparse
import json
import math
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, default=Path("../data"))
    parser.add_argument("--output-dir", type=Path, default=Path("outputs"))
    parser.add_argument("--min-count", type=int, default=20)
    args = parser.parse_args()
    data = json.loads((args.data_dir/"yelp-top300-final-0.7.json").read_text(encoding="utf-8-sig"))
    groups = defaultdict(list)
    for record in data:
        groups["-".join(record["ID"].split("-")[:2])].extend(record["quads"])
    workbook = load_workbook(args.data_dir/"overall-sentiment.xlsx", read_only=True, data_only=True)
    iterator = workbook.worksheets[0].values
    headers = next(iterator)
    wanted = ["ID", "business_name", "Positive", "Negative", "Neutral", "date"]
    positions = {h:headers.index(h) for h in wanted}
    meta = {}
    for values in iterator:
        key = values[positions["ID"]]
        if key in groups:
            assert key not in meta, ("Duplicate ID", key)
            meta[key] = {h:values[i] for h,i in positions.items()}
    workbook.close()
    assert len(meta) == len(groups)
    counts = Counter(meta[key]["business_name"] for key in groups)
    expected_ids = sorted((key for key in groups if counts[meta[key]["business_name"]] >= args.min_count), key=lambda key:tuple(map(int,key.split("-"))))
    aspect_names = set()
    for key in expected_ids:
        for quad in groups[key]:
            if quad[2] and quad[2].strip():
                aspect_names.add(quad[2].strip())
    final = load_workbook(args.output_dir/"yelp_data_processed.xlsx", read_only=True, data_only=True)
    rows = final.worksheets[0].values
    headers = list(next(rows))
    assert headers[:4] == ["ID", "business_name", "date", "sentiment"]
    assert set(headers[4:]) == aspect_names
    assert len(headers) == 4 + len(aspect_names)
    assert not any(h.startswith("quads-") or h in {"Positive", "Negative", "Neutral", "Date"} for h in headers)
    seen = []
    checked_quads = 0
    for values in rows:
        row = dict(zip(headers,values))
        key = row["ID"]
        seen.append(key)
        m = meta[key]
        assert row["business_name"] == m["business_name"], key
        assert row["date"] == datetime.fromisoformat(m["date"]), key
        labels = ["Positive", "Negative", "Neutral"]
        winner = max(labels,key=lambda label:m[label])
        expected_score = m[winner] * {"Positive":1,"Negative":-1,"Neutral":0}[winner]
        assert math.isclose(row["sentiment"],expected_score,abs_tol=1e-12), (key,"sentiment")
        vector = Counter()
        for n,quad in enumerate(groups[key],1):
            checked_quads += 1
            aspect = quad[2].strip() if quad[2] else ""
            if aspect:
                vector[aspect] += {"positive":1,"negative":-1,"negatives":-1,"neutral":0}[quad[3].strip().lower()]
        for aspect in aspect_names:
            assert row[aspect] == (vector[aspect] if aspect in vector else 999), (key,aspect)
    final.close()
    assert seen == expected_ids, "ID membership or ordering differs"
    intermediate = load_workbook(args.output_dir/"sentiment_counted.xlsx", read_only=True, data_only=True)
    intermediate_rows = intermediate.worksheets[0].values
    intermediate_headers = next(intermediate_rows)
    intermediate_ids = []
    for values in intermediate_rows:
        row = dict(zip(intermediate_headers, values))
        key = row["ID"]
        intermediate_ids.append(key)
        for n, quad in enumerate(groups[key], 1):
            for p, element in enumerate(quad, 1):
                assert row.get(f"quads-{n}-{p}") == (element or None), (key, n, p)
    intermediate.close()
    assert intermediate_ids == expected_ids
    result = {"status":"passed","source_sentences":len(data),"aggregated_reviews":len(groups),"verified_final_reviews":len(seen),"verified_quads":checked_quads,"aspect_columns":len(aspect_names),"final_columns":headers,"total_columns":len(headers),"retained_businesses":sum(v>=args.min_count for v in counts.values()),"removed_reviews":len(groups)-len(seen)}
    (args.output_dir/"independent_validation.json").write_text(json.dumps(result,indent=2),encoding="utf-8")
    print(json.dumps(result,indent=2))

if __name__ == "__main__":
    main()
