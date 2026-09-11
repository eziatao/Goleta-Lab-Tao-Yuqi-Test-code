"""Independently reconcile all three extension outputs against the input."""
import argparse
import json
import math
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

BASE = ["ID", "business_name", "date", "sentiment"]
DERIVED = ["consistency", "business_aspect_consistency", "average_sentiment_10"]


def read(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    iterator = workbook.worksheets[0].values
    headers = list(next(iterator))
    rows = [dict(zip(headers, values)) for values in iterator]
    workbook.close()
    return headers, rows


def equal(actual, expected, context):
    if isinstance(expected, (int, float)):
        assert actual is not None and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12), (context, actual, expected)
    else:
        assert actual == expected, (context, actual, expected)


def check_file(path, columns, expected_rows):
    workbook = load_workbook(path, read_only=True, data_only=True)
    iterator = workbook.worksheets[0].values
    assert list(next(iterator)) == columns, (path.name, "columns")
    count = 0
    for values in iterator:
        assert count < len(expected_rows), (path.name, "extra row")
        row = dict(zip(columns, values))
        for column in columns:
            equal(row.get(column), expected_rows[count].get(column), (path.name, count+2, column))
        count += 1
    workbook.close()
    assert count == len(expected_rows), (path.name, "missing row")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    columns, source = read(args.input)
    aspects = [c for c in columns if c not in BASE + DERIVED]
    counts = {a: sum(float(row[a]) != 999 for row in source) for a in aspects}
    selected = sorted(aspects, key=lambda a: (-counts[a], columns.index(a)))[:10]
    rows = [{c: row[c] for c in BASE + selected} for row in source]
    del source
    check_file(args.output_dir/"attribute_filter.xlsx", BASE+selected, rows)
    groups = defaultdict(list)
    for row in rows:
        values = [float(row[a]) for a in selected if float(row[a]) != 999]
        mean = math.fsum(values)/len(values) if values else None
        result = math.sqrt(math.fsum((v-mean)**2 for v in values)/len(values)) if values else None
        row["consistency"] = result
        if result is not None:
            groups[row["business_name"]].append(result)
    group_means = {name: math.fsum(values)/len(values) for name, values in groups.items()}
    for row in rows:
        row["business_aspect_consistency"] = group_means.get(row["business_name"])
    check_file(args.output_dir/"business_consistency_counted.xlsx", BASE+selected+DERIVED[:2], rows)
    rows.sort(key=lambda row: (tuple(map(int, row["ID"].split("-"))), row["ID"]))
    for i, row in enumerate(rows):
        previous = [float(rows[j]["sentiment"]) for j in range(max(0, i-10), i)]
        row["average_sentiment_10"] = math.fsum(previous)/len(previous) if previous else None
    check_file(args.output_dir/"yelp_data_processed_updated.xlsx", BASE+selected+DERIVED, rows)
    report = {"status": "passed", "verified_reviews": len(rows), "selected_counts": {a: counts[a] for a in selected},
              "output_columns": 4+len(selected)+3, "empty_consistency_reviews": sum(r["consistency"] is None for r in rows),
              "businesses_without_consistency": len({r["business_name"] for r in rows if r["business_aspect_consistency"] is None}),
              "verified_files": ["attribute_filter.xlsx", "business_consistency_counted.xlsx", "yelp_data_processed_updated.xlsx"]}
    (args.output_dir/"extension_validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
