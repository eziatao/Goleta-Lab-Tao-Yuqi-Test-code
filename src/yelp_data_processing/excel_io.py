"""Streaming Excel IO, text preservation and atomic replacement."""
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .common import DataError, Table


def read_excel(path, sheet=None, selected=None):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise DataError(f"Unknown worksheet: {sheet}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        iterator = ws.iter_rows(values_only=True)
        raw = list(next(iterator, ()))
        while raw and raw[-1] is None:
            raw.pop()
        if not raw or any(not isinstance(c, str) or not c.strip() for c in raw):
            raise DataError(f"{path}: expected non-empty text headers in row 1")
        headers = [c.strip() for c in raw]
        if len(headers) != len(set(headers)):
            raise DataError(f"{path}: duplicate headers")
        indices = [i for i, c in enumerate(headers) if selected is None or c in selected]
        columns = [headers[i] for i in indices]
        rows = []
        for values in iterator:
            if not any(v is not None for v in values):
                continue
            rows.append({headers[i]: values[i] for i in indices if i < len(values) and values[i] is not None})
        return Table(columns, rows)
    finally:
        wb.close()


def write_excel(table, path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise DataError("Excel output must use .xlsx")
    if not table.columns or len(table.columns) != len(set(table.columns)):
        raise DataError("Output headers must be non-empty and unique")
    if len(table.columns) > 16384 or len(table.rows) > 1048575:
        raise DataError("Output exceeds Excel's row/column limits")
    # Validate before opening the streaming writer: no truncation or partial output.
    for row in [dict.fromkeys(table.columns, None), *table.rows]:
        for value in list(row) + list(row.values()):
            if isinstance(value, str) and (len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value)):
                raise DataError("A cell contains illegal control characters or exceeds 32767 characters")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.freeze_panes = "B2"
    for i, col in enumerate(table.columns, 1):
        width = 36 if col == "business_name" else 24 if col in ("Date", "date") else max(16, min(34, len(col) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
    header = []
    for name in table.columns:
        cell = WriteOnlyCell(ws, name)
        cell.data_type = "s"
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234E70")
        cell.alignment = Alignment(vertical="center")
        header.append(cell)
    ws.row_dimensions[1].height = 24
    ws.append(header)
    for row in table.rows:
        cells = []
        for col in table.columns:
            value = row.get(col)
            if value is None:
                cells.append(None)
            elif isinstance(value, (str, date, datetime)):
                cell = WriteOnlyCell(ws, value)
                if isinstance(value, str):
                    cell.data_type = "s"  # '=...' remains literal text, never a formula.
                else:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                cells.append(cell)
            else:
                cells.append(value)
        ws.append(cells)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(table.columns))}{len(table.rows) + 1}"
    fd, temp = tempfile.mkstemp(suffix=".xlsx", prefix=".yelp-", dir=path.parent)
    os.close(fd)
    try:
        wb.save(temp)
        os.replace(temp, path)
    finally:
        wb.close()
        if os.path.exists(temp):
            os.unlink(temp)
    return path
